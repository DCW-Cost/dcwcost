"""Geometry QA for the deck.

LibreOffice cannot render in this container (it fails on a trivial deck too),
so instead of eyeballing images this measures the boxes directly: estimated
text height against declared box height, text-box overlaps, and safe-area
violations. Crude, but it catches the defect that actually matters — text
spilling out of its container.
"""
import math
import sys
from pptx import Presentation

EMU = 914400.0
SLIDE_W, SLIDE_H = 10.0, 5.625
SAFE = 0.35

# Montserrat is a fairly wide geometric sans; 0.56 em average is about right
# for mixed-case text. Erring wide means we flag borderline cases rather than
# miss them.
CHAR_EM = 0.56


def line_height_in(para, size_pt):
    ls = para.line_spacing
    if ls is None:
        return size_pt * 1.22 / 72.0
    if hasattr(ls, "pt"):          # a Length — pptxgenjs writes points
        return float(ls.pt) / 72.0
    return float(ls) * size_pt / 72.0  # a bare float is a multiplier


def para_size(para, default=12.0):
    for r in para.runs:
        if r.font.size:
            return r.font.size.pt
    if para.font.size:
        return para.font.size.pt
    return default


def estimate_height(tf, width_in):
    total = 0.0
    for para in tf.paragraphs:
        txt = "".join(r.text for r in para.runs)
        size = para_size(para)
        lh = line_height_in(para, size)
        if not txt.strip():
            total += lh
            continue
        # explicit newlines force breaks
        for seg in txt.split("\n"):
            cpl = max(1, int((width_in * 72) / (CHAR_EM * size)))
            total += max(1, math.ceil(len(seg) / cpl)) * lh
    return total


def main(path):
    p = Presentation(path)
    issues = []
    for idx, slide in enumerate(p.slides, 1):
        texts = []
        for sh in slide.shapes:
            if sh.left is None:
                continue
            x, y = sh.left / EMU, sh.top / EMU
            w, h = (sh.width or 0) / EMU, (sh.height or 0) / EMU

            if x < SAFE - 0.01 or x + w > SLIDE_W - SAFE + 0.02 or y < 0.15 or y + h > SLIDE_H - 0.12:
                issues.append(
                    f"S{idx:<2} SAFE-AREA  x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}"
                )

            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue

            need = estimate_height(sh.text_frame, w)
            label = sh.text_frame.text.replace("\n", " ")[:44]
            texts.append((x, y, w, max(h, need), label))
            if need > h * 1.05:
                issues.append(
                    f"S{idx:<2} OVERFLOW   need {need:.2f}\" have {h:.2f}\"  «{label}»"
                )

        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                ax, ay, aw, ah, at = texts[i]
                bx, by, bw, bh, bt = texts[j]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 0.08 and oy > 0.08:
                    issues.append(
                        f"S{idx:<2} COLLIDE    {oy:.2f}\" deep  «{at}» / «{bt}»"
                    )

    print(f"{len(issues)} issue(s)")
    for i in issues:
        print("  " + i)
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
